"""Rebuild reports/Ops_Analytics_Report.pptx. Run from the project root."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "reports" / "figures"
OUT = ROOT / "reports" / "Ops_Analytics_Report.pptx"

NAVY = RGBColor(0x1B, 0x4F, 0x72)
TEAL = RGBColor(0x14, 0x8F, 0x77)
ORANGE = RGBColor(0xD3, 0x54, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x2C, 0x3E, 0x50)
MUTED = RGBColor(0x5D, 0x6D, 0x7E)
CREAM = RGBColor(0xF7, 0xF5, 0xF2)
LINE = RGBColor(0xD5, 0xD8, 0xDC)

W, H = Inches(13.333), Inches(7.5)


def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def txt(slide, l, t, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return box


def footer(slide, page, total=11):
    rect(slide, 0, 7.28, 13.333, 0.22, NAVY)
    txt(slide, 0.4, 7.28, 8, 0.22, "Ops analytics  ·  not a live client", 9, False, WHITE)
    txt(slide, 11.4, 7.28, 1.5, 0.22, f"{page} / {total}", 9, False, WHITE, PP_ALIGN.RIGHT)


def pic(slide, name, l, t, w, h):
    slide.shapes.add_picture(str(FIG / name), Inches(l), Inches(t), Inches(w), Inches(h))


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1 title
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, NAVY)
    rect(s, 0, 0, 0.22, 7.5, ORANGE)
    txt(s, 0.7, 1.8, 12, 0.4, "OPS ANALYTICS", 14, True, ORANGE)
    txt(s, 0.7, 2.2, 12, 1.2, "Jobs reviewed, then a drop in weekly engagement", 32, True, WHITE)
    txt(s, 0.7, 4.0, 11, 0.8,
        "Two case studies. Pandas + SQLite. The job table has 8 rows;\nI still ran the queries. The product tables are real size.",
        16, False, WHITE)
    txt(s, 0.7, 6.4, 11, 0.3, "Practice project  ·  May–Aug 2014 events  ·  Nov 2020 jobs", 13, False, RGBColor(0xAE, 0xB6, 0xBF))

    # 2 ask
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, "What the brief asked", 22, True, WHITE)
    footer(s, 2)
    txt(s, 0.5, 1.0, 6, 0.35, "CASE 1  ·  JOB DATA", 13, True, ORANGE)
    for i, t in enumerate([
        "Jobs reviewed per hour per day (Nov 2020)",
        "Throughput (events / sec) and 7-day rolling",
        "Language share, last 30 days",
        "How to display duplicate rows",
    ]):
        txt(s, 0.5, 1.5 + i * 0.45, 6, 0.4, "•  " + t, 16, False, INK)
    txt(s, 7.0, 1.0, 5.8, 0.35, "CASE 2  ·  PRODUCT EVENTS", 13, True, ORANGE)
    for i, t in enumerate([
        "Weekly user engagement",
        "User growth",
        "Weekly retention of the signup cohort",
        "Weekly engagement per device",
        "Email engagement",
    ]):
        txt(s, 7.0, 1.5 + i * 0.45, 5.8, 0.4, "•  " + t, 16, False, INK)
    txt(s, 0.5, 4.4, 12, 2.2,
        "Engagement = unique users with at least one event_type = engagement that week.\n"
        "Weeks start Monday. IDs came in as floats; I cast them to int.\n"
        "Case 1 has no timestamp, only a date — so “per hour” is jobs / 24.",
        15, False, MUTED)

    # 3 case 1
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, "Case 1  ·  eight jobs, six days", 22, True, WHITE)
    footer(s, 3)
    pic(s, "jobs_per_day.png", 0.4, 1.0, 6.4, 2.9)
    pic(s, "language_share.png", 6.9, 1.0, 5.9, 2.8)
    txt(s, 0.5, 4.1, 12.3, 2.8,
        "Per hour: 0.0417 or 0.0833. Throughput: 0.000012–0.000023 events/sec. "
        "7-day rolling ends at 0.0000154 — and we only have 6 days, so it is a running mean, not a full window.\n\n"
        "Persian is 3 of 8 (37.5%). Everything else is one row.\n\n"
        "No exact duplicate rows. job_id 23 appears three times with different actors. "
        "I would show that with GROUP BY job_id HAVING COUNT(*) > 1, not SELECT DISTINCT.",
        15, False, INK)

    # 4 the drop
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, "Case 2  ·  weekly engagement falls in the week of 4 Aug", 22, True, WHITE)
    footer(s, 4)
    pic(s, "weekly_engagement.png", 0.35, 0.9, 8.3, 3.7)
    txt(s, 8.8, 1.1, 4.1, 0.3, "28 JUL  (PEAK)", 11, True, MUTED)
    txt(s, 8.8, 1.4, 4.1, 0.5, "1,443 users", 26, True, NAVY)
    txt(s, 8.8, 2.3, 4.1, 0.3, "4 AUG", 11, True, MUTED)
    txt(s, 8.8, 2.6, 4.1, 0.5, "1,266  (−12%)", 26, True, ORANGE)
    txt(s, 8.8, 3.5, 4.1, 0.3, "25 AUG", 11, True, MUTED)
    txt(s, 8.8, 3.8, 4.1, 0.5, "1,194  (−17%)", 26, True, ORANGE)
    txt(s, 0.5, 4.8, 12.3, 2.1,
        "340,832 events, 9,760 users in the event log, 1 May–31 Aug 2014. "
        "The week of 28 Apr is short (data starts Thursday 1 May). "
        "Events per engaged user also slip a little (14.9 → 13.5), so it is not only fewer people — they do slightly less.",
        15, False, INK)

    # 5 growth
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, "Growth is not the problem", 22, True, WHITE)
    footer(s, 5)
    pic(s, "user_growth.png", 0.35, 0.9, 7.8, 3.5)
    pic(s, "returning_vs_new.png", 8.2, 0.9, 4.8, 3.5)
    txt(s, 0.5, 4.6, 12.3, 2.3,
        "19,066 accounts created (Jan 2013–Aug 2014). 9,381 activated (49.2%). "
        "Weekly signups in August: 476 → 406 → 473 → 468 → 514. Activations follow.\n\n"
        "Returning engaged users: 1,153 (28 Jul) → 1,055 → 943 → 908. "
        "First-time engagers dip only in the week of 4 Aug (290 → 211) and then come back. "
        "The hole is existing users, not new ones.",
        15, False, INK)

    # 6 retention
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, "Signup-cohort retention", 22, True, WHITE)
    footer(s, 6)
    pic(s, "retention_by_age.png", 0.4, 1.0, 7.6, 3.8)
    txt(s, 8.3, 1.2, 4.6, 5.5,
        "7,298 signups inside the event window.\n\n"
        "Week 0  50.4%\nWeek 1  34.1%\nWeek 2  20.3%\nWeek 4  10.2%\nWeek 8   4.8%\n\n"
        "Week 0 is ~50% because pending accounts never engage. "
        "The brief said signup cohort, so I kept pending users in the denominator.",
        15, False, INK)

    # 7 device
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, "Phone and tablet fall harder than computer", 22, True, WHITE)
    footer(s, 7)
    pic(s, "device_weekly.png", 0.35, 0.9, 8.4, 3.8)
    txt(s, 8.9, 1.15, 4.0, 5.5,
        "Unique engaged users\n28 Jul → 25 Aug\n\n"
        "Computer   951 → 864   −9%\n"
        "Phone      589 → 441  −25%\n"
        "Tablet     250 → 163  −35%\n\n"
        "Families are a hand mapping of the 26 device strings (galaxy note → phone, “samsumg” tablet kept as tablet).",
        15, False, INK)

    # 8 email
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, "Email  ·  opens hold, clicks do not", 22, True, WHITE)
    footer(s, 8)
    pic(s, "email_rates.png", 0.35, 0.9, 8.4, 3.8)
    txt(s, 8.9, 1.1, 4.0, 5.6,
        "90,389 rows. Whole window:\nopen 33.6%  ·  CTR 14.8%\n\n"
        "28 Jul   open 35.2%  CTR 16.1%\n"
        "4 Aug    open 33.4%  CTR 10.8%\n"
        "25 Aug   open 35.0%  CTR 11.3%\n\n"
        "People still open the digest. Fewer click through in the same week engagement drops.",
        15, False, INK)

    # 9 so what
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, "What I would tell ops", 22, True, WHITE)
    footer(s, 9)
    items = [
        ("Not acquisition", "Signups and activations keep going up in August. New-to-engagement recovers after one week."),
        ("Existing users, mobile first", "Returning users fall and stay down. Phone −25%, tablet −35%, computer −9%."),
        ("Email CTA is a live lead", "Open rate is flat; CTR drops from 16.1% to 10.8% the week of 4 Aug."),
        ("I cannot name the bug from this file", "Event mix did not change. Next: mobile release log, digest template, August holidays by country."),
    ]
    for i, (h, b) in enumerate(items):
        y = 1.05 + i * 1.4
        rect(s, 0.5, y, 12.3, 1.25, WHITE)
        rect(s, 0.5, y, 0.12, 1.25, ORANGE if i < 3 else TEAL)
        txt(s, 0.9, y + 0.12, 11.6, 0.35, h, 18, True, NAVY)
        txt(s, 0.9, y + 0.52, 11.6, 0.6, b, 14, False, INK)

    # 10 method
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, "How this was run", 22, True, WHITE)
    footer(s, 10)
    txt(s, 0.6, 1.1, 12, 5.5,
        "Stack: Python (pandas) and SQLite. No extra database.\n\n"
        "Cleaning: drop empty Excel rows; cast user_id to int; parse timestamps; "
        "strip nothing else. Language and device strings are used as they arrived "
        "(including “portugese” and “samsumg galaxy tablet”).\n\n"
        "SQL lives in sql/01_job_ops.sql and sql/02_metric_spike.sql. "
        "The notebook loads the CSVs into an in-memory SQLite database and checks "
        "that the pandas and SQL answers match.\n\n"
        "I would not ship Case 1 as a dashboard. Eight rows. Case 2 is the actual investigation.",
        16, False, INK)

    # 11 close
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, NAVY)
    rect(s, 0, 0, 0.22, 7.5, ORANGE)
    txt(s, 0.7, 2.2, 12, 1.0, "1,443 → 1,266 engaged users\nin one week. Growth did not follow.", 28, True, WHITE)
    txt(s, 0.7, 4.4, 12, 1.2,
        "Check mobile clients and the digest click target first.\nDo not pause acquisition on this evidence.",
        16, False, RGBColor(0xD5, 0xD8, 0xDC))

    prs.save(OUT)
    print("wrote", OUT, OUT.stat().st_size)


if __name__ == "__main__":
    build()
