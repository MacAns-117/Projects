"""Rebuild Cosmetics_Report.pptx."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "reports/figures"
OUT = ROOT / "reports/Cosmetics_Report.pptx"
NAVY = RGBColor(0x1B, 0x4F, 0x72)
ORANGE = RGBColor(0xD3, 0x54, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x2C, 0x3E, 0x50)
CREAM = RGBColor(0xF7, 0xF5, 0xF2)
W, H = Inches(13.333), Inches(7.5)


def rect(s, l, t, w, h, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()


def txt(s, l, t, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = "Calibri"


def footer(s, n, total=8):
    rect(s, 0, 7.28, 13.333, 0.22, NAVY)
    txt(s, 0.4, 7.28, 10, 0.22, "Chemical components  ·  MedTourEasy practice project", 9, False, WHITE)
    txt(s, 11.4, 7.28, 1.5, 0.22, f"{n} / {total}", 9, False, WHITE, PP_ALIGN.RIGHT)


def pic(s, name, l, t, w, h):
    s.shapes.add_picture(str(FIG / name), Inches(l), Inches(t), Inches(w), Inches(h))


def header(s, title):
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, title, 22, True, WHITE)


def build():
    prs = Presentation(); prs.slide_width = W; prs.slide_height = H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, NAVY); rect(s, 0, 0, 0.22, 7.5, ORANGE)
    txt(s, 0.7, 2.0, 12, 0.3, "MEDTOUREASY  ·  1,472 SEPHORA SKUs", 14, True, ORANGE)
    txt(s, 0.7, 2.4, 12, 1.4, "Similar ingredients, cheaper jar", 32, True, WHITE)
    txt(s, 0.7, 4.1, 11, 1.0, "One-hot the formula list. t-SNE to look at it.\nCosine to pick a neighbor.", 18, False, WHITE)

    s = prs.slides.add_slide(blank); header(s, "Catalogue"); footer(s, 2)
    pic(s, "category_counts.png", 0.3, 0.9, 6.4, 3.4)
    pic(s, "price_by_category.png", 6.8, 0.9, 6.2, 3.4)
    txt(s, 0.5, 4.5, 12.3, 2.4,
        "Treatments are the expensive aisle (median $64.50). Cleansers are $28.\n"
        "Eye cream has the weakest mean rank (3.81). 116 brands; CLINIQUE leads SKU count (79).\n"
        "471 products have no skin-type flag — the 0/1 columns are incomplete.", 15, False, INK)

    s = prs.slides.add_slide(blank); header(s, "The matrix"); footer(s, 3)
    txt(s, 0.6, 1.1, 12, 5.5,
        "Filter: Label = Moisturizer AND Dry = 1  →  190 products.\n\n"
        "Split Ingredients on “, ”, lowercase, unique token → index. 2,233 tokens.\n"
        "decyl oleate sits at index 25 (the DataCamp check).\n\n"
        "Binary document-term matrix 190 × 2,233. Mean 35.1 ingredients / product. 98.4% zeros.\n\n"
        "t-SNE (2-D, learning_rate=200, random_state=42) for the map.\n"
        "Cosine on the raw binary rows for neighbors — that number is stable; t-SNE is not a metric.", 16, False, INK)

    s = prs.slides.add_slide(blank); header(s, "t-SNE map of dry moisturizers"); footer(s, 4)
    pic(s, "tsne_map.png", 0.35, 0.85, 7.6, 5.6)
    txt(s, 8.1, 1.2, 4.8, 5.2,
        "Axes are not ingredients.\nDistance ≈ shared formula.\n\n"
        "Orange: AmorePacific cushion $60 / 4.0\n"
        "Navy: Laneige BB cushion $38 / 4.3\n\n"
        "They sit next to each other.\nCosine 0.535 · 23 shared tokens.", 15, False, INK)

    s = prs.slides.add_slide(blank); header(s, "The $22 swap"); footer(s, 5)
    txt(s, 0.6, 1.1, 12, 5.5,
        "AmorePacific Color Control Cushion SPF 50+   $60   rank 4.0\n"
        "Laneige BB Cushion Hydra Radiance SPF 50     $38   rank 4.3\n\n"
        "Nearest product in the 190-row set (cosine 0.535).\n"
        "Next neighbor is only 0.333 — so this is not a crowded blob.\n\n"
        "Same job (tinted cushion + SPF), $22 less, slightly better rating.\n"
        "That is the whole point of the brief: read the back of the bottle without being a chemist.", 16, False, INK)

    s = prs.slides.add_slide(blank); header(s, "What this is not"); footer(s, 6)
    txt(s, 0.6, 1.1, 12, 5.5,
        "• Not a safety ranking. Shared silicones ≠ shared irritation profile.\n\n"
        "• Not “word embeddings.” It is one-hot presence. Order on the INCI list is ignored.\n\n"
        "• 120 rows say “Visit the boutique” instead of a formula. They contribute almost nothing.\n\n"
        "• t-SNE will jitter if you drop random_state. Use cosine for the actual recommendation.", 16, False, INK)

    s = prs.slides.add_slide(blank); header(s, "How it was run"); footer(s, 7)
    txt(s, 0.6, 1.1, 12, 5.5,
        "pandas + numpy one-hot, sklearn TSNE / cosine_similarity, SQLite for the catalogue, Excel for the tables.\n\n"
        "No Bokeh app. Static map in reports/figures/tsne_map.png.\n\n"
        "Notebook: notebooks/cosmetics_analysis.ipynb", 16, False, INK)

    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, NAVY); rect(s, 0, 0, 0.22, 7.5, ORANGE)
    txt(s, 0.7, 2.4, 12, 1.6, "Same silicones. Twenty-two dollars less.", 28, True, WHITE)

    prs.save(OUT)
    print("wrote", OUT, OUT.stat().st_size)


if __name__ == "__main__":
    build()
