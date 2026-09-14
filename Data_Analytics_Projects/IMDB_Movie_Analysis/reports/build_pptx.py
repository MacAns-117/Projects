"""Rebuild IMDB_Report.pptx."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "reports/figures"
OUT = ROOT / "reports/IMDB_Report.pptx"
NAVY = RGBColor(0x1B, 0x4F, 0x72)
ORANGE = RGBColor(0xD3, 0x54, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x2C, 0x3E, 0x50)
CREAM = RGBColor(0xF7, 0xF5, 0xF2)
W, H = Inches(13.333), Inches(7.5)


def rect(s, l, t, w, h, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()


def txt(s, l, t, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = "Calibri"


def footer(s, n, total=9):
    rect(s, 0, 7.28, 13.333, 0.22, NAVY)
    txt(s, 0.4, 7.28, 9, 0.22, "IMDb movies  ·  practice project", 9, False, WHITE)
    txt(s, 11.4, 7.28, 1.5, 0.22, f"{n} / {total}", 9, False, WHITE, PP_ALIGN.RIGHT)


def pic(s, name, l, t, w, h):
    s.shapes.add_picture(str(FIG / name), Inches(l), Inches(t), Inches(w), Inches(h))


def header(s, title):
    rect(s, 0, 0, 13.333, 7.5, CREAM)
    rect(s, 0, 0, 13.333, 0.7, NAVY)
    txt(s, 0.5, 0.18, 12, 0.4, title, 22, True, WHITE)


def build():
    prs = Presentation(); prs.slide_width = W; prs.slide_height = H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, NAVY); rect(s, 0, 0, 0.22, 7.5, ORANGE)
    txt(s, 0.7, 1.9, 12, 0.3, "IMDb  ·  4,916 unique titles  ·  1916–2016", 14, True, ORANGE)
    txt(s, 0.7, 2.3, 12, 1.4, "What actually moves the score?", 32, True, WHITE)
    txt(s, 0.7, 4.0, 11, 1.2, "Genre, runtime, and director beat budget for IMDb.\nBudget beats score for gross.", 18, False, WHITE)

    s = prs.slides.add_slide(blank); header(s, "A  ·  Genre"); footer(s, 2)
    pic(s, "genre_counts.png", 0.3, 0.9, 6.3, 4.0)
    pic(s, "genre_means.png", 6.7, 0.9, 6.3, 4.0)
    txt(s, 0.5, 5.1, 12.3, 1.9,
        "Drama is the most common tag (2,532). Documentary / biography / history / war sit ~7.1.\n"
        "Horror is 5.80. Comedy is everywhere (1,847) and still below the 6.44 all-movie mean.", 15, False, INK)

    s = prs.slides.add_slide(blank); header(s, "B  ·  Duration"); footer(s, 3)
    pic(s, "duration_scatter.png", 0.3, 0.9, 6.6, 4.3)
    pic(s, "duration_bins.png", 7.1, 0.9, 5.8, 3.5)
    txt(s, 0.5, 5.4, 12.3, 1.6,
        "r = 0.26. Under 90 min → 6.11 mean. Over 150 min → 7.44. Mode is 90 minutes.\n"
        "That is a mix/prestige effect, not a reason to pad the cut.", 15, False, INK)

    s = prs.slides.add_slide(blank); header(s, "C  ·  Language"); footer(s, 4)
    pic(s, "language_counts.png", 0.4, 1.0, 7.2, 4.0)
    txt(s, 7.8, 1.2, 5.0, 5.0,
        "English: 4,582 films, mean 6.39 (93% of the file).\n\n"
        "Japanese 7.35 (n=17)\nGerman 7.34 (n=19)\nFrench 7.04 (n=73)\n\n"
        "Small-n languages look better because this dump keeps the famous ones.", 15, False, INK)

    s = prs.slides.add_slide(blank); header(s, "D  ·  Directors (≥ 5 films)"); footer(s, 5)
    pic(s, "directors.png", 0.35, 0.9, 7.4, 4.2)
    txt(s, 7.9, 1.2, 5.0, 5.0,
        "214 directors with 5+ films here.\n\n"
        "Nolan 8.425 (8 films) — 100th percentile of that set.\n"
        "Tarantino 8.20 · Capra 8.06 · Kubrick 8.05\n\n"
        "90th percentile of means: 7.41\n"
        "One-film 9.5 TV rows are not this chart.", 15, False, INK)

    s = prs.slides.add_slide(blank); header(s, "E  ·  Budget, gross, profit"); footer(s, 6)
    pic(s, "budget_gross.png", 0.3, 0.9, 6.4, 4.2)
    pic(s, "top_profit.png", 6.85, 0.9, 6.1, 3.3)
    txt(s, 0.5, 5.3, 12.3, 1.7,
        "After dropping 10 local-currency budgets: r(budget, gross) = 0.63. r(budget, IMDb) ≈ 0.04.\n"
        "Avatar +$523.5M, Jurassic World +$502.2M, Titanic +$458.7M. Excel sheet E_correl has =CORREL().", 15, False, INK)

    s = prs.slides.add_slide(blank); header(s, "Five whys  ·  the duration pattern"); footer(s, 7)
    txt(s, 0.6, 1.0, 12.2, 5.8,
        "1. Why do longer films score higher? They sit with drama / biography / war, not 90-minute horror.\n\n"
        "2. Why those genres run long? Awards-circuit habit and editors keeping footage that works.\n\n"
        "3. Why the average moves: voters who finish 160 minutes are already bought in; short bins are padded with 5.x horror.\n\n"
        "4. Why this is not a lever: stretching a weak script does not mint Shawshank.\n\n"
        "5. Why it still matters: runtime is a proxy for intent. Budget tracks gross. IMDb tracks something else.", 16, False, INK)

    s = prs.slides.add_slide(blank); header(s, "What I would not tell a producer"); footer(s, 8)
    txt(s, 0.6, 1.1, 12, 5.5,
        "• Do not rank directors on a single 9.x.\n\n"
        "• Do not read Lady Vengeance as a $4.2B flop — that budget is not USD.\n\n"
        "• Do not treat English’s 6.39 as 'English films are worse.'\n\n"
        "• Do not use IMDb as a stand-in for profit. Avatar is both; plenty of 8.x films are not.", 16, False, INK)

    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, NAVY); rect(s, 0, 0, 0.22, 7.5, ORANGE)
    txt(s, 0.7, 2.2, 12, 1.8, "Spend money to chase gross.\nRuntime and genre describe the score.\nThey are not the same job.", 28, True, WHITE)

    prs.save(OUT)
    print("wrote", OUT, OUT.stat().st_size)


if __name__ == "__main__":
    build()
